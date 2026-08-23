"""Excel (.xlsx) rendering of GET /v1/sessions/{session_id}/treatment-plans.xlsx — one row per
accepted scenario that has a generated treatment plan.

Pure transform: no DB, no FastAPI. Takes the SAME TreatmentPlanStatus objects
get_treatment_plan() already returns (same trim, same staleness projection), so the Excel and
JSON views can never disagree — only presentation differs here.

That sourcing is load-bearing: get_treatment_plan's _normalize_plan_shape lifts legacy
(pre-v0.12) PlanJSON rows into the nested controls_to_be_implemented shape this module's
_row() unwraps. Feed it TreatmentPlanStatus objects from that GET — never hand-parsed
PlanJSON, which may still be the old flat-array shape.
"""
from __future__ import annotations

from typing import cast

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.api.schemas import TreatmentPlanStatus

_COLUMNS = (
    "session_id", "output_id", "plan_id",
    "threat_category", "threat_type", "threat_name", "threat_actors",
    "scenario_title", "scenario_statement", "risk_statement",
    "status", "treatment_strategy", "risk_level", "review_status", "risk_identification_date",
    "plan_title", "treatment_plan", "action_plan", "applicable_to_all_subsystems",
    "controls_to_be_implemented", "remediation_action_plan",
    "mitigation_timeline", "mitigation_owner", "risk_owner", "impacted_business_division",
    "error_message", "reason",
)

#: Multi-line/long-prose columns get wrap_text; everything else stays single-line.
_WRAP_COLUMNS = frozenset({
    "scenario_statement", "risk_statement", "action_plan",
    "controls_to_be_implemented", "remediation_action_plan",
})

#: Starting widths — narrow for ids/flags, wide for prose/list columns. Not load-bearing; a
#: reviewer can resize in Excel same as any spreadsheet.
_COLUMN_WIDTHS = {
    "session_id": 24, "output_id": 24, "plan_id": 24,
    "threat_category": 28, "threat_type": 28, "threat_name": 40, "threat_actors": 30,
    "scenario_title": 40,
    "scenario_statement": 50, "risk_statement": 50,
    "status": 12, "treatment_strategy": 14, "risk_level": 12, "review_status": 16,
    "risk_identification_date": 20,
    "plan_title": 30, "treatment_plan": 14, "action_plan": 50, "applicable_to_all_subsystems": 14,
    "controls_to_be_implemented": 50, "remediation_action_plan": 50,
    "mitigation_timeline": 30, "mitigation_owner": 20, "risk_owner": 20,
    "impacted_business_division": 28,
    "error_message": 30, "reason": 18,
}

#: Leading characters Excel/LibreOffice treat as a formula trigger — a cell value starting with
#: one of these executes as a formula on open, not shown as text. Plan content ultimately
#: originates from LLM output, so this is a real trust boundary, not a speculative guard (same
#: rule as app/api/results_excel.py).
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _sanitize(value):
    """Prefix a formula-triggering leading character with `'` so Excel renders it as literal
    text instead of executing it."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _controls_to_implement_cell(items: list[dict]) -> str:
    """One labelled block per gap-analysis entry, carrying every field the plan JSON holds —
    control_type, control_name, description, priority, control_code and control_library_id.

    The code/library id matter to a reviewer working the sheet: they are the join back to the
    control library, so a control can be looked up without re-opening the JSON response. Blocks
    are separated by a blank line; the column is wrap_text so each stays readable in one cell."""
    blocks = []
    for c in items:
        blocks.append(
            f"control_type: {c.get('control_type', '')}\n"
            f"control_name: {c.get('control_name', '')}\n"
            f"description: {c.get('description', '')}\n"
            f"priority: {c.get('priority', '')}\n"
            f"control_code: {c.get('control_code', '')}\n"
            f"control_library_id: {c.get('control_library_id', '')}"
        )
    return "\n\n".join(blocks)


def _remediation_actions_cell(items: list[dict]) -> str:
    """One '{action_id}: {action} (owner: ..., priority: ..., timeline: ...)' line per action."""
    lines = []
    for a in items:
        lines.append(f"{a.get('action_id', '')}: {a.get('action', '')} "
                     f"(owner: {a.get('owner', '')}, priority: {a.get('priority', '')}, "
                     f"timeline: {a.get('timeline', '')})")
    return "\n".join(lines)


def _row(session_id: str, status: TreatmentPlanStatus) -> dict:
    scenario = status.scenario or {}
    plan = status.plan or {}
    return {
        "session_id": session_id, "output_id": status.output_id, "plan_id": status.plan_id,
        "threat_category": scenario.get("threat_category"),
        "threat_type": scenario.get("threat_type"),
        "threat_name": scenario.get("threat_name"),
        "threat_actors": "; ".join(scenario.get("threat_actors") or []),
        "scenario_title": scenario.get("scenario_title"),
        "scenario_statement": scenario.get("scenario_statement"),
        "risk_statement": scenario.get("risk_statement"),
        "status": status.status, "treatment_strategy": status.treatment_strategy,
        "risk_level": status.risk_level, "review_status": status.review_status,
        "risk_identification_date": status.risk_identification_date,
        "plan_title": plan.get("title"),
        "treatment_plan": plan.get("treatment_plan"),
        "action_plan": plan.get("action_plan"),
        "applicable_to_all_subsystems": plan.get("applicable_to_all_subsystems"),
        "controls_to_be_implemented": _controls_to_implement_cell(
            (plan.get("controls_to_be_implemented") or {}).get("controls") or []),
        "remediation_action_plan": _remediation_actions_cell(
            plan.get("remediation_action_plan") or []),
        "mitigation_timeline": plan.get("mitigation_timeline"),
        "mitigation_owner": plan.get("mitigation_owner"),
        "risk_owner": plan.get("risk_owner"),
        "impacted_business_division": plan.get("impacted_business_division"),
        "error_message": status.error_message,
        "reason": status.reason,
    }


def build_treatment_plans_workbook(session_id: str, plans: list[TreatmentPlanStatus]) -> Workbook:
    """One 'Treatment Plans' sheet, one row per accepted scenario's generated plan."""
    wb = Workbook()
    # A freshly constructed Workbook always has exactly one active sheet; wb.active is
    # Optional only for a workbook whose sheets were all removed.
    ws: Worksheet = cast("Worksheet", wb.active)
    ws.title = "Treatment Plans"

    header_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    plain = Alignment(vertical="top")

    for col_idx, name in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = header_font
        ws.column_dimensions[get_column_letter(col_idx)].width = _COLUMN_WIDTHS[name]
    ws.freeze_panes = "A2"

    row_idx = 2
    for status in plans:
        row = _row(session_id, status)
        for col_idx, name in enumerate(_COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=_sanitize(row[name]))
            cell.alignment = wrap if name in _WRAP_COLUMNS else plain
        row_idx += 1

    return wb
